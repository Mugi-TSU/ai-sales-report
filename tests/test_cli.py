from openpyxl import load_workbook

from sales_report.cli import main


def test_mock_mode_creates_report(sample_csv, tmp_path, capsys):
    out = tmp_path / "out.xlsx"

    assert main([str(sample_csv), "-o", str(out), "--mock"]) == 0
    assert "サマリー" in load_workbook(out).sheetnames
    assert "2026-09" in capsys.readouterr().out


def test_missing_input_returns_error(tmp_path, capsys):
    assert main([str(tmp_path / "none.csv"), "--mock"]) == 1
    assert "見つかりません" in capsys.readouterr().err


def test_missing_api_key_returns_error(sample_csv, tmp_path, capsys, monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    monkeypatch.delenv("ANTHROPIC_AUTH_TOKEN", raising=False)

    assert main([str(sample_csv), "-o", str(tmp_path / "o.xlsx")]) == 1
    assert "ANTHROPIC_API_KEY" in capsys.readouterr().err


def test_ai_failure_still_writes_report_with_rule_based_comment(
    sample_csv, tmp_path, capsys, monkeypatch
):
    from sales_report import cli
    from sales_report.insights import InsightError

    class Failing:
        def write(self, summary):
            raise InsightError("AIの応答を解析できませんでした")

    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-ant-dummy")
    monkeypatch.setattr(cli, "_build_writer", lambda args: Failing())
    out = tmp_path / "o.xlsx"

    assert main([str(sample_csv), "-o", str(out)]) == 0
    assert out.exists()
    assert "簡易コメント" in capsys.readouterr().err
