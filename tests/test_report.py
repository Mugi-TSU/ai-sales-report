import zipfile
from pathlib import Path

from openpyxl import load_workbook

from sales_report.aggregate import summarize
from sales_report.insights import RuleBasedInsightWriter
from sales_report.loader import read_sales
from sales_report.report import write_report


def build(sample_csv: Path, tmp_path: Path) -> Path:
    summary = summarize(read_sales(sample_csv))
    out = tmp_path / "report.xlsx"
    write_report(summary, RuleBasedInsightWriter().write(summary), out)
    return out


def test_has_expected_sheets(sample_csv, tmp_path):
    wb = load_workbook(build(sample_csv, tmp_path))

    assert wb.sheetnames == ["サマリー", "店舗別", "カテゴリ別", "売れ筋"]


def test_summary_sheet_shows_latest_total(sample_csv, tmp_path):
    summary = summarize(read_sales(sample_csv))
    ws = load_workbook(build(sample_csv, tmp_path))["サマリー"]

    values = [c.value for row in ws.iter_rows() for c in row]
    assert summary.monthly_total[summary.latest_month] in values


def test_store_sheet_total_column_matches_overall(sample_csv, tmp_path):
    summary = summarize(read_sales(sample_csv))
    ws = load_workbook(build(sample_csv, tmp_path))["店舗別"]

    total_row = next(r for r in ws.iter_rows(values_only=True) if r[0] == "合計")
    assert list(total_row[1:]) == [summary.monthly_total[m] for m in summary.months]


def test_contains_charts(sample_csv, tmp_path):
    with zipfile.ZipFile(build(sample_csv, tmp_path)) as zf:
        charts = [n for n in zf.namelist() if n.startswith("xl/charts/chart")]

    assert len(charts) >= 2


def test_labels_comment_source(sample_csv, tmp_path):
    summary = summarize(read_sales(sample_csv))
    insight = RuleBasedInsightWriter().write(summary)
    ai_out, rule_out = tmp_path / "ai.xlsx", tmp_path / "rule.xlsx"

    write_report(summary, insight, ai_out, ai_generated=True)
    write_report(summary, insight, rule_out, ai_generated=False)

    assert load_workbook(ai_out)["サマリー"]["A6"].value == "AIの所見"
    assert load_workbook(rule_out)["サマリー"]["A6"].value == "所見（簡易）"
