"""集計結果と所見を Excel レポートに書き出す。"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from sales_report.models import Insight, MonthlyAmounts, SalesSummary

HEADER_FILL = PatternFill("solid", fgColor="3B2A20")
ACCENT_FILL = PatternFill("solid", fgColor="F3EDE4")
YEN_FORMAT = '#,##0"円"'
PCT_FORMAT = '+0.0"%";-0.0"%"'


def write_report(
    summary: SalesSummary, insight: Insight, path: Path, ai_generated: bool = False
) -> None:
    wb = Workbook()
    _write_summary_sheet(wb.active, summary, insight, ai_generated)
    _write_breakdown_sheet(
        wb.create_sheet("店舗別"), "店舗", summary.by_store, summary, LineChart
    )
    _write_breakdown_sheet(
        wb.create_sheet("カテゴリ別"),
        "カテゴリ",
        summary.by_category,
        summary,
        BarChart,
    )
    _write_top_products_sheet(wb.create_sheet("売れ筋"), summary)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _write_summary_sheet(
    ws: Worksheet, summary: SalesSummary, insight: Insight, ai_generated: bool
) -> None:
    ws.title = "サマリー"
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 90

    ws["A1"] = f"売上レポート（{summary.latest_month}）"
    ws["A1"].font = Font(size=16, bold=True)
    ws["A3"], ws["B3"] = "売上合計", summary.monthly_total[summary.latest_month]
    ws["B3"].number_format = YEN_FORMAT
    ws["A4"], ws["B4"] = "前月比", summary.mom_change_pct
    ws["B4"].number_format = PCT_FORMAT
    for cell in (ws["B3"], ws["B4"]):
        cell.font = Font(size=14, bold=True)
        cell.alignment = Alignment(horizontal="left")

    row = 6
    ws.cell(row, 1, "AIの所見" if ai_generated else "所見（簡易）").font = Font(
        bold=True
    )
    ws.cell(row, 2, insight.headline).font = Font(bold=True)
    for label, items in (("気づき", insight.findings), ("打ち手", insight.actions)):
        for i, text in enumerate(items):
            row += 1
            ws.cell(row, 1, label if i == 0 else None)
            ws.cell(row, 2, f"・{text}").alignment = Alignment(wrap_text=True)
    for r in range(6, row + 1):
        ws.cell(r, 1).fill = ACCENT_FILL


def _write_breakdown_sheet(
    ws: Worksheet,
    label: str,
    breakdown: dict[str, MonthlyAmounts],
    summary: SalesSummary,
    chart_type: type[LineChart] | type[BarChart],
) -> None:
    _write_header(ws, [label, *summary.months])
    for name, amounts in breakdown.items():
        ws.append([name, *(amounts[m] for m in summary.months)])
    ws.append(["合計", *(summary.monthly_total[m] for m in summary.months)])
    for row in ws.iter_rows(min_row=2, min_col=2):
        for cell in row:
            cell.number_format = YEN_FORMAT
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    ws.column_dimensions["A"].width = 16
    for col in range(2, len(summary.months) + 2):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = 14

    # 合計行を除いた各行を系列にして、月の推移を描く
    last_row = ws.max_row - 1
    chart = chart_type()
    chart.title = f"{label}別の売上推移"
    data = Reference(
        ws, min_col=1, max_col=len(summary.months) + 1, min_row=2, max_row=last_row
    )
    chart.add_data(data, from_rows=True, titles_from_data=True)
    chart.set_categories(
        Reference(ws, min_col=2, max_col=len(summary.months) + 1, min_row=1)
    )
    chart.width, chart.height = 18, 9
    ws.add_chart(chart, f"A{ws.max_row + 3}")


def _write_top_products_sheet(ws: Worksheet, summary: SalesSummary) -> None:
    _write_header(ws, ["順位", "商品", f"売上（{summary.latest_month}）"])
    for rank, (name, amount) in enumerate(summary.top_products, start=1):
        ws.append([rank, name, amount])
        ws.cell(ws.max_row, 3).number_format = YEN_FORMAT
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 18


def _write_header(ws: Worksheet, names: list[str]) -> None:
    ws.append(names)
    for cell in ws[ws.max_row]:
        cell.fill = HEADER_FILL
        cell.font = Font(bold=True, color="FFFFFF")
